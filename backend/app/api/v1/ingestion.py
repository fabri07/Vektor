"""
Ingestion pipeline endpoints.

POST   /ingestion/upload                — upload file, enqueue parsing job
GET    /ingestion/files                 — list files for current tenant
GET    /ingestion/files/{file_id}/preview   — get parsed_summary_json
POST   /ingestion/files/{file_id}/confirm  — confirm import (NEEDS_CONFIRMATION only)
"""

import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

import filetype
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.deps import get_current_tenant, get_current_user
from app.config.settings import get_settings
from app.integrations.s3 import S3Client
from app.jobs.ingestion_worker import (
    _FECHA_COLS,
    _GASTO_COLS,
    _PRODUCTO_COLS,
    _VENTA_COLS,
    _extract_amounts_from_text,
    _analyze_headers,
    _rows_to_dicts,
    process_image_ocr,
    process_spreadsheet,
    process_text_document,
)
from app.main import limiter
from app.observability.logger import get_logger
from app.persistence.db.session import get_db_session
from app.persistence.models.file import (
    PROCESSING_STATUS_DONE,
    PROCESSING_STATUS_FAILED,
    PROCESSING_STATUS_NEEDS_CONFIRMATION,
    PROCESSING_STATUS_PENDING,
    PROCESSING_STATUS_PROCESSING,
    UploadedFile,
)
from app.persistence.models.tenant import Tenant
from app.persistence.models.user import User
from app.persistence.repositories.file_repository import FileRepository
from app.schemas.ingestion import (
    ConfirmIngestionRequest,
    ConfirmIngestionResponse,
    FilePreviewResponse,
    FileStatusItem,
    UploadResponse,
)

router = APIRouter()

logger = get_logger(__name__)

# ── Constants ─────────────────────────────────────────────────────────────────

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

# MIME types detected by filetype (magic bytes)
_SPREADSHEET_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",  # xlsx
}
_TEXT_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",  # docx
}
_IMAGE_MIMES = {
    "image/jpeg",
    "image/png",
    "image/heif",  # filetype returns image/heif for .heic
    "image/heic",
}

# Extensions that filetype can't detect (plain text) — fallback
_PLAIN_TEXT_EXTENSIONS = {".txt", ".csv"}

ALLOWED_MIMES = _SPREADSHEET_MIMES | _TEXT_MIMES | _IMAGE_MIMES | {"text/plain", "text/csv"}

FileHint = Literal["ventas", "gastos", "stock", "general"]

# Filename sanitization: keep alphanumerics, dots, dashes, underscores only
_SAFE_FILENAME_RE = re.compile(r"[^a-zA-Z0-9.\-_]")


# ── Helpers ───────────────────────────────────────────────────────────────────

def _sanitize_filename(filename: str) -> str:
    """Remove path traversal and unsafe characters from a filename."""
    # Strip directory components
    filename = filename.replace("\\", "/").split("/")[-1]
    # Remove unsafe characters
    filename = _SAFE_FILENAME_RE.sub("_", filename)
    # Ensure non-empty
    return filename or "upload"


def _detect_mime(content: bytes, filename: str) -> str:
    """
    Detect real MIME type from magic bytes via filetype.
    Falls back to extension-based detection for plain-text formats
    (CSV, TXT) that filetype cannot identify from binary signatures.
    Raises HTTPException 415 if the type is not allowed.
    """
    kind = filetype.guess(content[:2048])
    if kind is not None:
        detected = kind.mime
    else:
        ext = f".{filename.rsplit('.', 1)[-1].lower()}" if "." in filename else ""
        if ext in _PLAIN_TEXT_EXTENSIONS:
            detected = "text/csv" if ext == ".csv" else "text/plain"
        else:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail="Tipo de archivo no reconocido. Tipos aceptados: xlsx, csv, txt, docx, jpg, png, heic.",
            )

    if detected not in ALLOWED_MIMES:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Tipo de archivo '{detected}' no soportado.",
        )
    return detected


def _pick_job(mime: str) -> object:
    """Return the Celery task to enqueue for a given MIME type."""
    if mime in _IMAGE_MIMES:
        return process_image_ocr
    if mime in _SPREADSHEET_MIMES or mime in ("text/csv",):
        return process_spreadsheet
    return process_text_document


# ── Ingestion data insertion helpers ─────────────────────────────────────────

def _parse_amount(raw: Any) -> Decimal | None:
    if raw is None:
        return None
    s = re.sub(r"[$\s]", "", str(raw).strip())
    if not s:
        return None
    # Argentine/European format: 1.500,50 → 1500.50
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        val = Decimal(s)
        return val if val > 0 else None
    except InvalidOperation:
        return None


def _parse_date(raw: Any) -> date | None:
    if raw is None:
        return None
    s = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _find_col(headers: list[str], keywords: set[str]) -> str | None:
    """Return first header whose normalized name contains any of the keywords."""
    for h in headers:
        norm = h.lower().strip().replace(" ", "_")
        if any(k in norm for k in keywords):
            return h
    return None


_NOMBRE_COLS: set[str] = {"producto", "descripcion", "descripción", "nombre", "articulo", "artículo", "item", "name", "concepto", "detalle"}
_PRECIO_VENTA_COLS: set[str] = {"precio_venta", "precio", "price", "p_venta"}
_COSTO_COLS: set[str] = {"costo", "cost", "precio_costo", "p_costo"}
_STOCK_COLS: set[str] = {"stock", "cantidad", "inventario", "units", "qty", "existencia"}
_SKU_COLS: set[str] = {"sku", "codigo", "código", "code", "ref", "id_producto"}


async def _insert_confirmed_data(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    summary: dict[str, Any],
    confirmed_fields: dict[str, bool],
) -> dict[str, int]:
    """Parse parsed_summary_json and insert rows into sales/expense/product tables."""
    from app.persistence.models.product import Product  # noqa: PLC0415
    from app.persistence.models.transaction import ExpenseEntry, SaleEntry  # noqa: PLC0415
    from sqlalchemy import select  # noqa: PLC0415

    today = date.today()
    counts: dict[str, int] = {"ventas": 0, "gastos": 0, "productos": 0}
    file_type = summary.get("file_type", "spreadsheet")

    if file_type == "spreadsheet":
        rows: list[dict[str, Any]] = summary.get("ventas_detectadas", [])
        if not rows:
            return counts

        headers = list(rows[0].keys())

        fecha_col = _find_col(headers, _FECHA_COLS)
        venta_col = _find_col(headers, _VENTA_COLS)
        gasto_col = _find_col(headers, _GASTO_COLS)
        nombre_col = _find_col(headers, _NOMBRE_COLS)
        precio_col = _find_col(headers, _PRECIO_VENTA_COLS)
        costo_col = _find_col(headers, _COSTO_COLS)
        stock_col = _find_col(headers, _STOCK_COLS)
        sku_col = _find_col(headers, _SKU_COLS)

        wants_ventas = bool(confirmed_fields.get("ventas") and summary.get("has_venta") and venta_col)
        wants_gastos = bool(confirmed_fields.get("gastos") and summary.get("has_gasto") and gasto_col)
        wants_productos = bool(confirmed_fields.get("productos") and summary.get("has_producto") and nombre_col)

        for row in rows:
            tx_date = (_parse_date(row.get(fecha_col)) if fecha_col else None) or today

            if wants_ventas:
                amount = _parse_amount(row.get(venta_col))
                if amount:
                    session.add(SaleEntry(
                        tenant_id=tenant_id,
                        amount=amount,
                        quantity=1,
                        transaction_date=tx_date,
                        payment_method="cash",
                        notes="Importado desde archivo",
                    ))
                    counts["ventas"] += 1

            if wants_gastos:
                amount = _parse_amount(row.get(gasto_col))
                if amount:
                    desc_raw = row.get(nombre_col) if nombre_col else None
                    desc = str(desc_raw).strip()[:499] if desc_raw and str(desc_raw).strip() not in {"None", "nan", ""} else "Gasto importado"
                    session.add(ExpenseEntry(
                        tenant_id=tenant_id,
                        amount=amount,
                        category="importado",
                        transaction_date=tx_date,
                        description=desc,
                        payment_method="transfer",
                    ))
                    counts["gastos"] += 1

        if wants_productos:
            for row in rows:
                name = str(row.get(nombre_col, "")).strip()[:299]
                if not name or name.lower() in {"none", "nan", ""}:
                    continue
                price = _parse_amount(row.get(precio_col)) if precio_col else None
                cost = _parse_amount(row.get(costo_col)) if costo_col else None
                try:
                    stock_raw = row.get(stock_col) if stock_col else None
                    stock_val = int(float(str(stock_raw))) if stock_raw not in (None, "", "None", "nan") else 0
                except (ValueError, TypeError):
                    stock_val = 0
                sku_raw = row.get(sku_col) if sku_col else None
                sku = str(sku_raw).strip()[:99] if sku_raw and str(sku_raw).strip() not in {"", "None", "nan"} else None

                result = await session.execute(
                    select(Product).where(
                        Product.tenant_id == tenant_id,
                        Product.name == name,
                    )
                )
                existing = result.scalar_one_or_none()
                if existing:
                    if price:
                        existing.sale_price_ars = price
                    if cost:
                        existing.unit_cost_ars = cost
                    if stock_val > 0:
                        existing.stock_units = stock_val
                    if sku:
                        existing.sku = sku
                else:
                    session.add(Product(
                        tenant_id=tenant_id,
                        name=name,
                        sku=sku,
                        sale_price_ars=price or Decimal("0"),
                        unit_cost_ars=cost,
                        stock_units=stock_val,
                    ))
                counts["productos"] += 1

    else:
        # text / image: lists of {linea, montos}
        if confirmed_fields.get("ventas"):
            for entry in summary.get("ventas_detectadas", []):
                for m in entry.get("montos", []):
                    amount = _parse_amount(m)
                    if amount:
                        session.add(SaleEntry(
                            tenant_id=tenant_id,
                            amount=amount,
                            quantity=1,
                            transaction_date=today,
                            payment_method="cash",
                            notes=str(entry.get("linea", ""))[:499],
                        ))
                        counts["ventas"] += 1

        if confirmed_fields.get("gastos"):
            for entry in summary.get("gastos_detectados", []):
                for m in entry.get("montos", []):
                    amount = _parse_amount(m)
                    if amount:
                        session.add(ExpenseEntry(
                            tenant_id=tenant_id,
                            amount=amount,
                            category="importado",
                            transaction_date=today,
                            description=str(entry.get("linea", ""))[:499] or "Gasto importado",
                            payment_method="transfer",
                        ))
                        counts["gastos"] += 1

    await session.flush()
    return counts


# ── Sync fallback (beta: Celery/Redis unavailable) ───────────────────────────


async def _process_file_sync(
    record: UploadedFile,
    session: AsyncSession,
) -> None:
    """Process a file synchronously when Celery is unavailable.

    Reuses the parsing helpers from ingestion_worker but runs inside the
    existing request session instead of creating a separate Celery-owned
    engine.  On failure the file is marked FAILED so the user sees a clear
    status instead of being stuck in PENDING forever.
    """
    import csv as _csv  # noqa: PLC0415
    import io  # noqa: PLC0415

    repo = FileRepository(session)
    try:
        record.processing_status = PROCESSING_STATUS_PROCESSING
        await repo.save(record)
        await session.flush()

        s3 = S3Client()
        content = await s3.download(record.s3_key)
        mime = record.content_type

        summary: dict[str, Any] = {"warnings": []}

        if mime in _IMAGE_MIMES:
            summary["file_type"] = "image"
            summary["confidence"] = "LOW"
            try:
                import pytesseract  # noqa: PLC0415
                from PIL import Image, UnidentifiedImageError  # noqa: PLC0415

                try:
                    img = Image.open(io.BytesIO(content))
                    raw_text = pytesseract.image_to_string(img, lang="spa+eng")
                except UnidentifiedImageError:
                    raw_text = ""
                    summary["warnings"].append(
                        "No se pudo abrir la imagen (formato no soportado por Pillow)."
                    )

                amounts = _extract_amounts_from_text(raw_text)
                summary.update({"raw_text_preview": raw_text[:500], **amounts})
            except ImportError:
                summary["error"] = "OCR no disponible en este entorno"
                summary["ventas_detectadas"] = []
                summary["gastos_detectados"] = []
                summary["stock_detectado"] = []

        elif mime in _SPREADSHEET_MIMES or mime in ("text/csv",):
            summary["file_type"] = "spreadsheet"
            is_csv = mime == "text/csv" or record.original_filename.endswith(".csv")
            if is_csv:
                text = content.decode("utf-8", errors="replace")
                reader = _csv.reader(io.StringIO(text))
                rows = list(reader)
                if not rows:
                    summary.update({"confidence": "MEDIUM", "ventas_detectadas": [], "rows_processed": 0})
                else:
                    headers = rows[0]
                    data_rows = rows[1:]
                    analysis = _analyze_headers(headers)
                    summary.update(analysis)
                    summary["headers"] = headers
                    summary["rows_processed"] = len(data_rows)
                    summary["ventas_detectadas"] = _rows_to_dicts(headers, data_rows[:50])
            else:
                import openpyxl  # noqa: PLC0415

                wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                ws = wb.active
                all_rows = list(ws.iter_rows(values_only=True))  # type: ignore[union-attr]
                if not all_rows:
                    summary.update({"confidence": "MEDIUM", "ventas_detectadas": [], "rows_processed": 0})
                else:
                    headers = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(all_rows[0])]
                    data_rows_raw = [list(r) for r in all_rows[1:]]
                    analysis = _analyze_headers(headers)
                    summary.update(analysis)
                    summary["headers"] = headers
                    summary["rows_processed"] = len(data_rows_raw)
                    summary["ventas_detectadas"] = _rows_to_dicts(headers, data_rows_raw[:50])
                wb.close()

        else:
            # text / docx
            summary["file_type"] = "text"
            summary["confidence"] = "MEDIUM"
            is_docx = (
                mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                or record.original_filename.endswith(".docx")
            )
            if is_docx:
                import docx  # noqa: PLC0415

                doc = docx.Document(io.BytesIO(content))
                raw_text = "\n".join(p.text for p in doc.paragraphs)
            else:
                raw_text = content.decode("utf-8", errors="replace")

            amounts = _extract_amounts_from_text(raw_text)
            summary.update({"raw_text_preview": raw_text[:500], **amounts})

        record.parsed_summary_json = summary
        record.processing_status = PROCESSING_STATUS_NEEDS_CONFIRMATION
        await repo.save(record)

        logger.info(
            "ingestion.sync_fallback.done",
            file_id=str(record.id),
            file_type=summary.get("file_type"),
            confidence=summary.get("confidence"),
        )

    except Exception as exc:
        logger.error(
            "ingestion.sync_fallback.failed",
            file_id=str(record.id),
            error=str(exc),
        )
        record.parsed_summary_json = {"error": str(exc)}
        record.processing_status = PROCESSING_STATUS_FAILED
        await repo.save(record)


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post(
    "/upload",
    response_model=UploadResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Upload a file and enqueue ingestion job",
)
@limiter.limit("20/hour")
async def upload_file(
    request: Request,
    file: UploadFile = File(...),
    file_hint: FileHint = Query(default="general"),
    tenant: Tenant = Depends(get_current_tenant),
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_db_session),
) -> UploadResponse:
    content = await file.read()

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="El archivo supera el tamaño máximo de 10 MB.",
        )

    filename = _sanitize_filename(file.filename or "upload")
    detected_mime = _detect_mime(content, filename)

    # Build S3 key: uploads/{tenant_id}/{uuid}/{filename}
    file_uuid = uuid.uuid4()
    s3_key = f"uploads/{tenant.tenant_id}/{file_uuid}/{filename}"

    s3 = S3Client()
    stored_key = await s3.upload_to_key(content=content, key=s3_key, content_type=detected_mime)

    record = UploadedFile(
        tenant_id=tenant.tenant_id,
        uploaded_by=current_user.user_id,
        original_filename=filename,
        s3_key=stored_key,
        content_type=detected_mime,
        size_bytes=len(content),
        purpose=file_hint,
        status="uploaded",
        processing_status=PROCESSING_STATUS_PENDING,
    )
    repo = FileRepository(session)
    saved = await repo.save(record)

    if get_settings().USE_LOCAL_FALLBACK:
        await _process_file_sync(saved, session)
        return UploadResponse(file_id=saved.id, status="PROCESSING")

    # Enqueue parsing job — fall back to sync processing if Celery/Redis
    # is unavailable (beta: single Railway service without workers).
    job = _pick_job(detected_mime)
    try:
        job.delay(str(saved.id), str(tenant.tenant_id))  # type: ignore[attr-defined]
    except Exception:
        logger.warning(
            "ingestion.celery_unavailable",
            file_id=str(saved.id),
            msg="Celery/Redis no disponible, procesando archivo de forma síncrona.",
        )
        await _process_file_sync(saved, session)
        return UploadResponse(file_id=saved.id, status="PROCESSING")

    return UploadResponse(file_id=saved.id, status="PROCESSING")


@router.get(
    "/files",
    response_model=list[FileStatusItem],
    summary="List ingested files for the current tenant",
)
async def list_files(
    processing_status: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
    tenant: Tenant = Depends(get_current_tenant),
    session: AsyncSession = Depends(get_db_session),
) -> list[UploadedFile]:
    repo = FileRepository(session)
    return await repo.list_by_tenant_filtered(
        tenant_id=tenant.tenant_id,
        processing_status=processing_status,
        limit=limit,
        offset=offset,
    )


@router.get(
    "/files/{file_id}/preview",
    response_model=FilePreviewResponse,
    summary="Get parsed summary for user review",
)
async def get_file_preview(
    file_id: uuid.UUID,
    tenant: Tenant = Depends(get_current_tenant),
    session: AsyncSession = Depends(get_db_session),
) -> FilePreviewResponse:
    repo = FileRepository(session)
    record = await repo.get_by_id(file_id, tenant.tenant_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")

    if record.processing_status in (PROCESSING_STATUS_PENDING, "PROCESSING"):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"El archivo aún se está procesando (estado: {record.processing_status}).",
        )

    return FilePreviewResponse(
        file_id=record.id,
        processing_status=record.processing_status,
        parsed_summary_json=record.parsed_summary_json,
    )


@router.delete(
    "/files/{file_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Delete an uploaded file",
)
async def delete_file(
    file_id: uuid.UUID,
    tenant: Tenant = Depends(get_current_tenant),
    session: AsyncSession = Depends(get_db_session),
) -> None:
    repo = FileRepository(session)
    record = await repo.get_by_id(file_id, tenant.tenant_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")

    try:
        s3 = S3Client()
        await s3.delete(record.s3_key)
    except Exception as exc:
        logger.warning("ingestion.delete.s3_failed", file_id=str(file_id), error=str(exc))

    await repo.delete(record)
    await session.commit()


@router.post(
    "/files/{file_id}/reprocess",
    status_code=status.HTTP_202_ACCEPTED,
    summary="Re-enqueue a PENDING or FAILED file for processing",
)
async def reprocess_file(
    file_id: uuid.UUID,
    tenant: Tenant = Depends(get_current_tenant),
    session: AsyncSession = Depends(get_db_session),
) -> dict:
    repo = FileRepository(session)
    record = await repo.get_by_id(file_id, tenant.tenant_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")

    if record.processing_status not in (PROCESSING_STATUS_PENDING, PROCESSING_STATUS_FAILED):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"El archivo no puede reprocesarse (estado actual: {record.processing_status}).",
        )

    record.processing_status = PROCESSING_STATUS_PENDING
    record.parsed_summary_json = None
    await repo.save(record)

    if get_settings().USE_LOCAL_FALLBACK:
        await _process_file_sync(record, session)
    else:
        job = _pick_job(record.content_type)
        try:
            job.delay(str(record.id), str(tenant.tenant_id))
        except Exception:
            await _process_file_sync(record, session)

    return {"file_id": str(file_id), "status": "requeued"}


@router.post(
    "/files/{file_id}/confirm",
    response_model=ConfirmIngestionResponse,
    summary="Confirm ingestion of parsed data",
)
async def confirm_file(
    file_id: uuid.UUID,
    body: ConfirmIngestionRequest,
    tenant: Tenant = Depends(get_current_tenant),
    session: AsyncSession = Depends(get_db_session),
) -> ConfirmIngestionResponse:
    repo = FileRepository(session)
    record = await repo.get_by_id(file_id, tenant.tenant_id)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado.")

    if record.processing_status != PROCESSING_STATUS_NEEDS_CONFIRMATION:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"El archivo no está pendiente de confirmación "
                f"(estado actual: {record.processing_status})."
            ),
        )

    # Insert parsed rows into business tables, then mark done
    updated_summary = dict(record.parsed_summary_json or {})
    updated_summary["confirmed_fields"] = body.confirmed_fields

    counts = await _insert_confirmed_data(session, tenant.tenant_id, updated_summary, body.confirmed_fields)
    updated_summary["imported_counts"] = counts

    record.parsed_summary_json = updated_summary
    record.processing_status = PROCESSING_STATUS_DONE
    await repo.save(record)

    logger.info(
        "ingestion.confirm.done",
        file_id=str(file_id),
        ventas=counts["ventas"],
        gastos=counts["gastos"],
        productos=counts["productos"],
    )

    # Enqueue score recalculation — BSL will aggregate newly confirmed data
    from app.application.services.score_trigger_service import (  # noqa: PLC0415
        trigger_score_recalculation,
    )

    try:
        trigger_score_recalculation.delay(str(tenant.tenant_id), str(file_id))
    except Exception:
        logger.warning("ingestion.confirm.score_trigger_failed", file_id=str(file_id))

    parts: list[str] = []
    if counts["ventas"]:
        parts.append(f"{counts['ventas']} venta(s)")
    if counts["gastos"]:
        parts.append(f"{counts['gastos']} gasto(s)")
    if counts["productos"]:
        parts.append(f"{counts['productos']} producto(s)")

    message = (
        f"Importados: {', '.join(parts)}. La puntuación será recalculada."
        if parts
        else "Datos confirmados. La puntuación de salud será recalculada."
    )

    return ConfirmIngestionResponse(
        file_id=record.id,
        status=PROCESSING_STATUS_DONE,
        message=message,
    )
