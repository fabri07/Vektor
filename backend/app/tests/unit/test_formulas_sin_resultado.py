"""E6c-2 — celda vacía, fórmula sin resultado y resultado cacheado son tres cosas.

El defecto, medido antes de arreglarlo: ``openpyxl`` con ``data_only=True``
devuelve ``None`` tanto para una celda vacía como para una fórmula cuyo resultado
el archivo nunca guardó. Un ``.xlsx`` generado por script —o uno que Excel nunca
recalculó— entra con **todas sus columnas calculadas leídas como vacías**.

Y "vacío" es un dato válido, no un error: la fila no falla, no avisa, no queda en
"Otros". Simplemente el total nunca existió. Es la peor forma de pérdida: la que
no deja rastro.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook

from app.application.services.file_parsing import (
    marcar_formulas_sin_resultado,
    parse_uploaded_content,
)
from app.domain.ingestion_limits import (
    FORMULA_SIN_RESULTADO,
    LIMITES,
    LimitesDeArchivo,
    RelojDeParsing,
    es_formula_sin_resultado,
)

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _libro_con_los_tres_casos() -> bytes:
    """Una fila por caso, para que el test pueda distinguirlos.

    ``openpyxl`` escribe la fórmula SIN resultado cacheado (no evalúa nada), que
    es exactamente el archivo que genera cualquier script.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["fecha", "detalle", "monto"])
    ws.append(["2024-03-05", "con formula sin calcular", "=10+20"])
    ws.append(["2024-03-06", "con valor cargado", 500])
    ws.append(["2024-03-07", "con la celda vacia", None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _hojas_crudas(contenido: bytes) -> dict[str, list[list[Any]]]:
    import openpyxl

    libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.active
    return {hoja.title: [list(f) for f in hoja.iter_rows(values_only=True)]}


def test_los_tres_estados_quedan_distinguidos() -> None:
    """El corazón de la entrega: sin esto, dos de los tres son el mismo ``None``."""
    contenido = _libro_con_los_tres_casos()
    hojas = _hojas_crudas(contenido)

    # Antes de marcar: la fórmula y la celda vacía son indistinguibles.
    assert hojas["Ventas"][1][2] is None
    assert hojas["Ventas"][3][2] is None

    marcadas, verificado = marcar_formulas_sin_resultado(
        contenido, hojas, RelojDeParsing(LIMITES), LIMITES
    )
    assert verificado
    assert marcadas == 1

    assert es_formula_sin_resultado(hojas["Ventas"][1][2]), "la fórmula sin calcular"
    assert hojas["Ventas"][2][2] == 500, "el valor cargado no se toca"
    assert hojas["Ventas"][3][2] is None, "la celda vacía SIGUE vacía"


def test_una_formula_con_resultado_cacheado_no_se_marca() -> None:
    """El caso normal —Excel guardó el resultado— tiene que seguir igual.

    Si se marcara, un archivo perfectamente bueno mandaría filas a revisión.
    """
    wb = Workbook()
    ws = wb.active
    ws.append(["monto"])
    ws.append([123])  # como si la fórmula ya estuviera resuelta
    buffer = io.BytesIO()
    wb.save(buffer)
    contenido = buffer.getvalue()

    hojas = _hojas_crudas(contenido)
    marcadas, verificado = marcar_formulas_sin_resultado(
        contenido, hojas, RelojDeParsing(LIMITES), LIMITES
    )
    assert (marcadas, verificado) == (0, True)


def test_sin_celdas_vacias_no_se_paga_la_segunda_lectura() -> None:
    """La verificación cuesta una segunda pasada: sólo se paga si hay algo que
    distinguir. Un archivo sin ningún ``None`` no tiene fórmulas sin resultado
    por definición."""
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    buffer = io.BytesIO()
    wb.save(buffer)

    hojas = _hojas_crudas(buffer.getvalue())
    assert marcar_formulas_sin_resultado(
        buffer.getvalue(), hojas, RelojDeParsing(LIMITES), LIMITES
    ) == (0, True)


def test_cuando_no_se_puede_verificar_se_dice() -> None:
    """Un cero que en realidad significa "no miré" se lee como "no hay fórmulas
    rotas". Son dos respuestas distintas y la diferencia tiene que viajar."""
    contenido = _libro_con_los_tres_casos()
    hojas = _hojas_crudas(contenido)
    # Presupuesto agotado: la verificación no llega a correr entera.
    limites = LimitesDeArchivo(segundos=0.0)
    marcadas, verificado = marcar_formulas_sin_resultado(
        contenido, hojas, RelojDeParsing(limites), limites
    )
    assert verificado is False, (
        "sin poder verificar hay que declararlo, no devolver 0 como si estuviera limpio"
    )
    assert marcadas == 0


def test_un_archivo_roto_no_rompe_el_parseo() -> None:
    """No poder abrir el libro para verificar es un resultado, no un fallo: el
    archivo ya se leyó, y perder el parseo entero por la verificación sería
    cambiar una pérdida silenciosa por una ruidosa."""
    assert marcar_formulas_sin_resultado(
        b"esto no es un xlsx", {"x": [[None]]}, RelojDeParsing(LIMITES), LIMITES
    ) == (0, False)


def test_el_summary_reporta_las_formulas() -> None:
    """El conteo viaja al summary, que es lo que lee el confirm para avisar."""
    summary = parse_uploaded_content(_libro_con_los_tres_casos(), _XLSX, "x.xlsx")
    assert summary["formulas_sin_resultado"] == 1
    assert summary["formulas_verificadas"] is True


def test_la_fila_con_la_formula_no_entra_como_si_estuviera_vacia() -> None:
    """La prueba que da sentido a todo lo anterior.

    La marca tiene que llegar hasta el motivo que ve el usuario: la fila va a
    "Otros" diciendo que hay una fórmula sin calcular, no "falta el monto" — que
    mandaría a mapear una columna que ya está mapeada— ni, peor, entrar con la
    celda leída como vacía.

    **Nunca se sustituye por cero ni por vacío**: ése era el comportamiento previo
    y es el que hacía la pérdida invisible.
    """
    from app.application.services.ingestion_import_service import _label_fila_sin_monto

    etiqueta = _label_fila_sin_monto(FORMULA_SIN_RESULTADO)
    assert "fórmula" in etiqueta
    # Y dice cómo resolverlo, no sólo qué pasó.
    assert "recalcule" in etiqueta or "CSV" in etiqueta
    # La celda realmente vacía conserva su mensaje de siempre.
    assert "fórmula" not in _label_fila_sin_monto(None)
